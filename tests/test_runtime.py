import io
import os
import unittest
import requests
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from openpyxl import Workbook

import app as live
from liveyours import documents
from liveyours.ai.providers import _safe_provider_error


class RuntimeTests(unittest.TestCase):
    def test_provider_error_diagnostics_are_short_and_redact_credentials(self):
        response = SimpleNamespace(
            json=lambda: {"error": {"code": "rate_limit", "message": "Rate limit exceeded for key=sk-or-v1-1234567890abcdef"}},
            text="",
        )
        details = _safe_provider_error(response)
        self.assertIn("code=rate_limit", details)
        self.assertIn("Rate limit exceeded", details)
        self.assertNotIn("sk-or-v1-1234567890abcdef", details)

    def setUp(self):
        self.original = {
            "APP_USERNAME": live.APP_USERNAME,
            "APP_PASSWORD": live.APP_PASSWORD,
            "CLOUD_AI_FALLBACK": live.CLOUD_AI_FALLBACK,
            "OPENROUTER_API_KEY": live.OPENROUTER_API_KEY,
            "BYTEZ_API_KEY": live.BYTEZ_API_KEY,
            "LLM7_API_KEY": live.LLM7_API_KEY,
            "OMNIROUTE_URL": live.OMNIROUTE_URL,
            "FREELLMAPI_API_KEY": live.FREELLMAPI_API_KEY,
            "FREELLMAPI_URL": live.FREELLMAPI_URL,
        }
        live.APP_USERNAME = "liveyours"
        live.APP_PASSWORD = "test-only-password"
        live.CLOUD_AI_FALLBACK = False
        live.FREELLMAPI_API_KEY = ""
        live.FREELLMAPI_URL = "http://127.0.0.1:3001/v1/chat/completions"
        live.clear_login_failures()
        self.client = live.app.test_client()
        self.auth = ("liveyours", "test-only-password")

    def tearDown(self):
        for name, value in self.original.items():
            setattr(live, name, value)

    def test_dashboard_requires_auth_and_loads_for_authenticated_user(self):
        with patch.object(live, "send_from_directory", return_value=live.Response("LiveYours shell")):
            response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"LiveYours", response.data)

    def test_private_api_requires_auth_without_browser_challenge(self):
        response = self.client.get("/api/tasks")
        self.assertEqual(response.status_code, 401)
        self.assertNotIn("WWW-Authenticate", response.headers)

    def test_login_sets_http_only_session_and_authenticates_session_endpoint(self):
        unauthenticated = self.client.get("/api/auth/session")
        self.assertEqual(unauthenticated.status_code, 200)
        self.assertEqual(unauthenticated.json, {"authenticated": False})

        response = self.client.post(
            "/api/auth/login",
            headers={"Origin": "http://localhost"},
            json={"username": "liveyours", "password": "test-only-password"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json["success"])
        self.assertIn("HttpOnly", response.headers["Set-Cookie"])
        self.assertIn("SameSite=Lax", response.headers["Set-Cookie"])
        self.assertEqual(self.client.get("/api/auth/session").json, {"authenticated": True})

    def test_invalid_login_does_not_create_authenticated_session(self):
        response = self.client.post(
            "/api/auth/login",
            headers={"Origin": "http://localhost"},
            json={"username": "liveyours", "password": "incorrect"},
        )
        self.assertEqual(response.status_code, 401)
        self.assertEqual(self.client.get("/api/auth/session").json, {"authenticated": False})

    def test_login_throttles_repeated_bad_credentials(self):
        for _ in range(live.LOGIN_FAILURE_LIMIT):
            response = self.client.post(
                "/api/auth/login",
                headers={"Origin": "http://localhost"},
                json={"username": "liveyours", "password": "incorrect"},
            )
            self.assertEqual(response.status_code, 401)

        response = self.client.post(
            "/api/auth/login",
            headers={"Origin": "http://localhost"},
            json={"username": "liveyours", "password": "test-only-password"},
        )
        self.assertEqual(response.status_code, 429)
        self.assertEqual(response.headers["Retry-After"], str(live.LOGIN_FAILURE_WINDOW_SECONDS))

    def test_basic_auth_failures_share_login_throttle(self):
        for _ in range(live.LOGIN_FAILURE_LIMIT):
            response = self.client.get("/api/tasks", auth=("liveyours", "incorrect"))
            self.assertEqual(response.status_code, 401)

        response = self.client.post(
            "/api/auth/login",
            headers={"Origin": "http://localhost"},
            json={"username": "liveyours", "password": "test-only-password"},
        )
        self.assertEqual(response.status_code, 429)

    def test_successful_login_clears_failed_attempts(self):
        for _ in range(live.LOGIN_FAILURE_LIMIT - 1):
            self.client.post(
                "/api/auth/login",
                headers={"Origin": "http://localhost"},
                json={"username": "liveyours", "password": "incorrect"},
            )

        response = self.client.post(
            "/api/auth/login",
            headers={"Origin": "http://localhost"},
            json={"username": "liveyours", "password": "test-only-password"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(live._login_failures, [])

    def test_login_rejects_missing_origin(self):
        response = self.client.post(
            "/api/auth/login",
            json={"username": "liveyours", "password": "test-only-password"},
        )
        self.assertEqual(response.status_code, 403)

    def test_privacy_endpoint_reports_active_mode(self):
        response = self.client.get("/api/privacy", auth=self.auth)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"cloud_ai_fallback": False, "ai_mode": "local-only"})

    def test_mutations_reject_missing_origin(self):
        response = self.client.post("/api/not-a-route", auth=self.auth, json={})
        self.assertEqual(response.status_code, 403)

    def test_finance_activity_and_trip_reject_non_finite_numbers(self):
        invalid_requests = [
            ("/api/finance/add", {"amount": "NaN", "type": "expense"}),
            ("/api/finance/balance", {"amount": "Infinity"}),
            ("/api/activities/add", {"type": "Run", "distance": "NaN", "duration_minutes": 30}),
            ("/api/plan_trip", {"destination": "Bizerte", "budget": "NaN"}),
        ]
        with patch.object(live, "get_db") as get_db:
            for endpoint, payload in invalid_requests:
                with self.subTest(endpoint=endpoint):
                    response = self.client.post(endpoint, auth=self.auth, headers={"Origin": "http://localhost"}, json=payload)
                    self.assertEqual(response.status_code, 400)
        get_db.assert_not_called()

    def test_sleep_rejects_quality_outside_one_to_ten(self):
        with patch.object(live, "get_db") as get_db:
            response = self.client.post(
                "/api/sleep", auth=self.auth, headers={"Origin": "http://localhost"}, json={"hours": 8, "quality": 11}
            )
        self.assertEqual(response.status_code, 400)
        get_db.assert_not_called()

    def test_objective_create_accepts_month_and_saves_month_start(self):
        database = MagicMock()
        with patch.object(live, "get_db", return_value=database):
            response = self.client.post(
                "/api/tasks", auth=self.auth, headers={"Origin": "http://localhost"},
                json={"text": "Finish capstone", "category": "Project", "target_month": "2026-11"},
            )
        self.assertEqual(response.status_code, 200)
        cursor = database.__enter__.return_value.cursor.return_value.__enter__.return_value
        self.assertIn("target_month", cursor.execute.call_args.args[0])
        self.assertEqual(cursor.execute.call_args.args[1][-1].isoformat(), "2026-11-01")

    def test_objective_create_rejects_invalid_month_before_database_access(self):
        with patch.object(live, "get_db") as get_db:
            response = self.client.post(
                "/api/tasks", auth=self.auth, headers={"Origin": "http://localhost"},
                json={"text": "Finish capstone", "target_month": "2026-13"},
            )
        self.assertEqual(response.status_code, 400)
        get_db.assert_not_called()

    def test_objective_update_can_move_objective_to_another_month(self):
        database = MagicMock()
        with patch.object(live, "get_db", return_value=database):
            response = self.client.post(
                "/api/tasks/update", auth=self.auth, headers={"Origin": "http://localhost"},
                json={"id": 7, "text": "Finish capstone", "category": "Project", "description": "", "target_month": "2027-01"},
            )
        self.assertEqual(response.status_code, 200)
        cursor = database.__enter__.return_value.cursor.return_value.__enter__.return_value
        self.assertIn("target_month = COALESCE", cursor.execute.call_args.args[0])
        self.assertEqual(cursor.execute.call_args.args[1][3].isoformat(), "2027-01-01")

    def test_health_check_queries_database(self):
        with patch.object(live, "get_db") as get_db:
            get_db.return_value.__enter__.return_value.cursor.return_value.__enter__.return_value.execute.return_value = None
            response = self.client.get("/healthz")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"status": "ok"})

    def test_trip_planner_calculates_trip_and_calls_ai(self):
        completion = {"choices": [{"message": {"content": "A budget itinerary."}}]}
        with patch.object(live, "request_ai", return_value=completion) as request_ai:
            response = self.client.post(
                "/api/plan_trip",
                auth=self.auth,
                headers={"Origin": "http://localhost"},
                json={"destination": "Bizerte", "budget": 100},
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["trip_plan"], "A budget itinerary.")
        self.assertGreater(response.json["fuel_cost_tnd"], 0)
        request_ai.assert_called_once()

    def test_trip_planner_returns_calculated_outline_when_ai_is_slow(self):
        with patch.object(live, "request_ai", side_effect=requests.Timeout("upstream timeout")) as request_ai:
            response = self.client.post(
                "/api/plan_trip",
                auth=self.auth,
                headers={"Origin": "http://localhost"},
                json={"destination": "Bizerte", "budget": 100},
            )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json["ai_generated"])
        self.assertIn("AI suggestions are temporarily unavailable", response.json["notice"])
        self.assertIn("Leave Tunis", response.json["trip_plan"])
        self.assertIn("Bizerte", response.json["trip_plan"])
        self.assertGreater(response.json["fuel_cost_tnd"], 0)
        self.assertEqual(request_ai.call_args.kwargs["timeout"], 20)

    def test_trip_planner_fallback_uses_configured_home_city_after_provider_error(self):
        with patch.dict(os.environ, {"HOME_CITY": "Tunis"}), patch.object(
            live, "request_ai", side_effect=requests.HTTPError("provider unavailable")
        ):
            response = self.client.post(
                "/api/plan_trip",
                auth=self.auth,
                headers={"Origin": "http://localhost"},
                json={"destination": "Bizerte", "budget": 100},
            )

        self.assertEqual(response.status_code, 200)
        self.assertIn("Leave Tunis", response.json["trip_plan"])
        self.assertNotIn("Leave Bizerte", response.json["trip_plan"])
        self.assertIn("AI suggestions are temporarily unavailable", response.json["notice"])

    def test_render_does_not_try_pc_localhost_when_cloud_fallback_is_disabled(self):
        live.OMNIROUTE_URL = "http://127.0.0.1:20128/v1/chat/completions"
        live.CLOUD_AI_FALLBACK = False
        with patch.dict(os.environ, {"RENDER": "true"}):
            with self.assertRaises(RuntimeError):
                live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]})
            self.assertIn("Hosted AI is disabled", live.ai_unavailable_message())

    def test_render_uses_openrouter_as_cloud_fallback(self):
        live.OMNIROUTE_URL = "https://stale-temporary-tunnel.example/v1/chat/completions"
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-only-key"
        response = type(
            "Response",
            (),
            {"raise_for_status": lambda self: None, "json": lambda self: {"choices": [{"message": {"content": "OK"}}]}},
        )()
        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(requests, "post", return_value=response) as post:
            result = live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]})
        self.assertEqual(result["choices"][0]["message"]["content"], "OK")
        self.assertEqual(post.call_args.args[0], live.OPENROUTER_URL)
        self.assertEqual(post.call_args.kwargs["json"]["model"], "openrouter/free")
        self.assertEqual(post.call_args.kwargs["json"]["max_tokens"], 1024)

    def test_openrouter_preserves_larger_completion_budget(self):
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-openrouter-key"
        response = type(
            "Response",
            (),
            {"raise_for_status": lambda self: None, "json": lambda self: {"choices": [{"message": {"content": "OK"}}]}},
        )()
        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(requests, "post", return_value=response) as post:
            live.request_ai({"model": "quick", "max_tokens": 1500, "messages": [{"role": "user", "content": "test"}]})
        self.assertEqual(post.call_args.kwargs["json"]["max_tokens"], 1500)

    def test_render_skips_omniroute_even_when_a_remote_tunnel_is_configured(self):
        live.OMNIROUTE_URL = "https://stale-temporary-tunnel.example/v1/chat/completions"
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-only-key"
        response = type(
            "Response",
            (),
            {"raise_for_status": lambda self: None, "json": lambda self: {"choices": [{"message": {"content": "OK"}}]}},
        )()
        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(requests, "post", return_value=response) as post:
            live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]})
        self.assertEqual(post.call_count, 1)
        self.assertEqual(post.call_args.args[0], live.OPENROUTER_URL)

    def test_local_falls_back_from_omniroute_to_freellmapi(self):
        live.FREELLMAPI_API_KEY = "test-unified-key"
        unavailable = type(
            "Response",
            (),
            {"raise_for_status": lambda self: (_ for _ in ()).throw(requests.ConnectionError("offline"))},
        )()
        successful = type(
            "Response",
            (),
            {"raise_for_status": lambda self: None, "json": lambda self: {"choices": [{"message": {"content": "Local routed answer"}}]}},
        )()

        with patch.dict(os.environ, {"RENDER": "false"}, clear=False), patch.object(
            requests, "post", side_effect=[unavailable, successful]
        ) as post:
            result = live.request_ai({"model": "study", "max_tokens": 200, "messages": [{"role": "user", "content": "test"}]})

        self.assertEqual(result["choices"][0]["message"]["content"], "Local routed answer")
        self.assertEqual(post.call_count, 2)
        self.assertEqual(post.call_args_list[1].args[0], "http://127.0.0.1:3001/v1/chat/completions")
        self.assertEqual(post.call_args_list[1].kwargs["json"]["model"], "auto")
        self.assertEqual(post.call_args_list[1].kwargs["headers"]["Authorization"], "Bearer test-unified-key")

    def test_render_never_contacts_local_freellmapi(self):
        live.FREELLMAPI_API_KEY = "test-unified-key"
        live.CLOUD_AI_FALLBACK = False
        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(requests, "post") as post:
            with self.assertRaises(RuntimeError):
                live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]})
        post.assert_not_called()

    def test_render_falls_through_when_provider_returns_empty_completion(self):
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-openrouter-key"
        live.BYTEZ_API_KEY = "test-bytez-key"
        empty_response = type(
            "Response",
            (),
            {
                "raise_for_status": lambda self: None,
                "json": lambda self: {"choices": [{"message": {"content": "  "}}]},
            },
        )()
        good_response = type(
            "Response",
            (),
            {
                "raise_for_status": lambda self: None,
                "json": lambda self: {"choices": [{"message": {"content": "Useful answer"}}]},
            },
        )()

        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(
            requests, "post", side_effect=[empty_response, good_response]
        ) as post:
            result = live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]})

        self.assertEqual(result["choices"][0]["message"]["content"], "Useful answer")
        self.assertEqual(post.call_count, 2)
        self.assertEqual(post.call_args_list[0].args[0], live.OPENROUTER_URL)
        self.assertEqual(post.call_args_list[1].args[0], live.BYTEZ_URL)

    def test_render_normalizes_openai_content_parts_to_text(self):
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-openrouter-key"
        response = type(
            "Response",
            (),
            {"raise_for_status": lambda self: None, "json": lambda self: {"choices": [{"message": {
                "content": [{"type": "text", "text": "Useful "}, {"type": "text", "text": "answer"}],
            }}]}},
        )()
        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(requests, "post", return_value=response):
            result = live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]})
        self.assertEqual(result["choices"][0]["message"]["content"], "Useful answer")

    def test_render_falls_through_to_bytez_when_openrouter_fails(self):
        live.OMNIROUTE_URL = "https://stale-temporary-tunnel.example/v1/chat/completions"
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-openrouter-key"
        live.BYTEZ_API_KEY = "test-bytez-key"

        failed_response = type(
            "Response",
            (),
            {
                "raise_for_status": lambda self: (_ for _ in ()).throw(
                    requests.HTTPError("unauthorized", response=type("ErrorResponse", (), {"status_code": 401})())
                )
            },
        )()
        successful_response = type(
            "Response",
            (),
            {"raise_for_status": lambda self: None, "json": lambda self: {"choices": [{"message": {"content": "OK"}}]}},
        )()

        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(
            requests, "post", side_effect=[failed_response, successful_response]
        ) as post:
            result = live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]})

        self.assertEqual(result["choices"][0]["message"]["content"], "OK")
        self.assertEqual(post.call_count, 2)
        self.assertEqual(post.call_args_list[0].args[0], live.OPENROUTER_URL)
        self.assertEqual(post.call_args_list[1].args[0], live.BYTEZ_URL)
        self.assertEqual(post.call_args_list[1].kwargs["headers"]["Authorization"], "test-bytez-key")
        self.assertEqual(post.call_args_list[1].kwargs["json"]["model"], "Qwen/Qwen3-4B")

    def test_render_uses_llm7_after_openrouter_and_bytez_fail(self):
        live.OMNIROUTE_URL = "https://stale-temporary-tunnel.example/v1/chat/completions"
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-openrouter-key"
        live.BYTEZ_API_KEY = "test-bytez-key"
        live.LLM7_API_KEY = "test-llm7-key"

        failed_response = type(
            "Response",
            (),
            {
                "raise_for_status": lambda self: (_ for _ in ()).throw(
                    requests.HTTPError("unavailable", response=type("ErrorResponse", (), {"status_code": 503})())
                )
            },
        )()
        successful_response = type(
            "Response",
            (),
            {
                "raise_for_status": lambda self: None,
                "json": lambda self: {"choices": [{"message": {"content": "OK"}}]},
            },
        )()

        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(
            requests, "post", side_effect=[failed_response, failed_response, successful_response]
        ) as post:
            result = live.request_ai({"model": "study", "messages": [{"role": "user", "content": "test"}]}, timeout=12)

        self.assertEqual(result["choices"][0]["message"]["content"], "OK")
        self.assertEqual(post.call_count, 3)
        self.assertEqual(post.call_args_list[2].args[0], live.LLM7_URL)
        self.assertEqual(post.call_args_list[2].kwargs["headers"]["Authorization"], "Bearer test-llm7-key")
        self.assertEqual(post.call_args_list[2].kwargs["json"]["model"], live.LLM7_MODEL)
        self.assertLessEqual(sum(post.call_args_list[0].kwargs["timeout"]), 4.2)

    def test_cloud_fallback_budget_is_shared_across_provider_attempts(self):
        live.OMNIROUTE_URL = "http://127.0.0.1:20128/v1/chat/completions"
        live.CLOUD_AI_FALLBACK = True
        live.OPENROUTER_API_KEY = "test-openrouter-key"
        live.BYTEZ_API_KEY = "test-bytez-key"
        live.LLM7_API_KEY = "test-llm7-key"
        failure = requests.Timeout("slow provider")
        with patch.dict(os.environ, {"RENDER": "true"}), patch.object(
            requests, "post", side_effect=[failure, failure, type("Response", (), {
                "raise_for_status": lambda self: None,
                "json": lambda self: {"choices": [{"message": {"content": "OK"}}]},
            })()]
        ) as post:
            live.request_ai({"model": "quick", "messages": [{"role": "user", "content": "test"}]}, timeout=12)

        budgets = [sum(call.kwargs["timeout"]) for call in post.call_args_list]
        self.assertEqual(len(budgets), 3)
        self.assertTrue(all(budget <= 4.2 for budget in budgets), repr(budgets))

    def test_upload_text_storage_is_capped(self):
        page = SimpleNamespace(extract_text=lambda: "A" * (live.MAX_STORED_DOCUMENT_CHARS + 500))
        with patch.object(documents, "PdfReader", return_value=SimpleNamespace(pages=[page])):
            text, truncated = live.extract_document_text(io.BytesIO(b"%PDF"), "notes.pdf")
        self.assertEqual(len(text), live.MAX_STORED_DOCUMENT_CHARS)
        self.assertTrue(truncated)

    def test_library_extracts_bounded_csv_rows(self):
        text, truncated = live.extract_document_text(
            io.BytesIO("Topic,Hours\nMath,2\nSecurity,3\n".encode()), "revision.csv"
        )
        self.assertEqual(text, "Topic\tHours\nMath\t2\nSecurity\t3")
        self.assertFalse(truncated)

    def test_library_upload_accepts_csv_and_xlsx(self):
        workbook = Workbook()
        workbook.active.append(["Topic", "Hours"])
        workbook.active.append(["Security", 3])
        xlsx = io.BytesIO()
        workbook.save(xlsx)
        samples = [
            ("revision.csv", b"Topic,Hours\nSecurity,3\n"),
            ("revision.xlsx", xlsx.getvalue()),
        ]
        db = MagicMock()
        db.__enter__.return_value = db

        with patch.object(live, "get_db", return_value=db) as get_db:
            for filename, payload in samples:
                with self.subTest(filename=filename):
                    response = self.client.post(
                        "/api/upload",
                        auth=self.auth,
                        headers={"Origin": "http://localhost"},
                        data={"files": (io.BytesIO(payload), filename)},
                        content_type="multipart/form-data",
                    )
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(response.json["files"], [filename])
        self.assertEqual(get_db.call_count, 2)

    def test_library_extracts_xlsx_sheet_names_and_values(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Revision plan"
        sheet.append(["Topic", "Hours"])
        sheet.append(["Security", 3])
        data = io.BytesIO()
        workbook.save(data)

        text, truncated = live.extract_document_text(io.BytesIO(data.getvalue()), "revision.xlsx")

        self.assertIn("Sheet: Revision plan", text)
        self.assertIn("Topic\tHours", text)
        self.assertIn("Security\t3", text)
        self.assertFalse(truncated)

    def test_xlsx_extraction_respects_character_budget(self):
        workbook = Workbook()
        workbook.active.append(["A" * 100])
        data = io.BytesIO()
        workbook.save(data)
        with patch.object(live, "MAX_STORED_DOCUMENT_CHARS", 20):
            text, truncated = live.extract_document_text(io.BytesIO(data.getvalue()), "large.xlsx")
        self.assertLessEqual(len(text), 20)
        self.assertTrue(truncated)

    def test_library_upload_reports_when_extracted_text_is_capped(self):
        page = SimpleNamespace(extract_text=lambda: "study material " * 3_000)
        db = MagicMock()
        db.__enter__.return_value = db
        cursor = MagicMock()
        cursor.__enter__.return_value = cursor
        db.cursor.return_value = cursor
        with patch.object(documents, "PdfReader", return_value=SimpleNamespace(pages=[page])), patch.object(live, "get_db", return_value=db):
            response = self.client.post(
                "/api/upload",
                auth=self.auth,
                headers={"Origin": "http://localhost"},
                data={"files": (io.BytesIO(b"%PDF"), "course.pdf")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["files"], ["course.pdf"])
        self.assertIn("safe document-processing limits", response.json["notices"][0])
        self.assertEqual(cursor.execute.call_args.args[1][1], ("study material " * 3_000).strip()[:live.MAX_STORED_DOCUMENT_CHARS])

    def test_upload_rejects_oversized_images_before_ocr(self):
        image = SimpleNamespace(format="PNG", width=5001, height=5000)
        with patch.object(documents.Image, "open", return_value=image), patch.object(documents.pytesseract, "image_to_string") as ocr:
            with self.assertRaisesRegex(ValueError, "13 megapixels"):
                live.extract_document_text(io.BytesIO(b"image"), "scan.png")
        ocr.assert_not_called()

    def test_upload_rejects_too_many_pdf_pages(self):
        pages = [SimpleNamespace(extract_text=lambda: "text") for _ in range(live.MAX_DOCUMENT_PAGES + 1)]
        with patch.object(documents, "PdfReader", return_value=SimpleNamespace(pages=pages)):
            with self.assertRaisesRegex(ValueError, "50 pages"):
                live.extract_document_text(io.BytesIO(b"%PDF"), "book.pdf")

    def test_chat_document_excerpt_is_relevant_and_bounded(self):
        document = ("opening material " * 120) + ("middlehit relevant security policy " * 100) + ("closing material " * 120)
        excerpt = live.select_document_excerpt(document, "explain middlehit security policy", 2400)
        self.assertLessEqual(len(excerpt), 2400)
        self.assertIn("middlehit", excerpt)

    def test_flashcard_source_is_bounded_and_samples_later_sections(self):
        document = "START " + ("x" * 16_990) + " MIDDLE " + ("y" * 12_990) + " END"
        excerpt, sampled = live.sample_document_text(document, live.MAX_FLASHCARD_SOURCE_CHARS)
        self.assertLessEqual(len(excerpt), live.MAX_FLASHCARD_SOURCE_CHARS)
        self.assertTrue(sampled)
        self.assertIn("START", excerpt)
        self.assertIn("MIDDLE", excerpt)
        self.assertIn("END", excerpt)

    def test_flashcards_report_sampling_for_long_document(self):
        db = MagicMock()
        db.__enter__.return_value = db
        cursor = MagicMock()
        cursor.__enter__.return_value = cursor
        cursor.fetchone.return_value = {"filename": "course.pdf", "content": "A" * 50_000}
        db.cursor.return_value = cursor
        response_payload = {"choices": [{"message": {"content": '[{"q":"Q","a":"A"}]'}}]}
        with patch.object(live, "get_db", return_value=db), patch.object(live, "request_ai", return_value=response_payload) as request_ai:
            response = self.client.post("/api/flashcards", auth=self.auth, headers={"Origin": "http://localhost"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json["flashcards"]), 1)
        self.assertIn("representative sections", response.json["source_note"])
        prompt = request_ai.call_args.args[0]["messages"][0]["content"]
        self.assertLess(len(prompt), live.MAX_FLASHCARD_SOURCE_CHARS + 500)

    def test_chat_sends_bounded_recent_history_to_ai(self):
        def fake_db(rows=None, new_id=None):
            db = MagicMock()
            db.__enter__.return_value = db
            cursor = MagicMock()
            cursor.__enter__.return_value = cursor
            cursor.fetchone.return_value = {"id": new_id} if new_id is not None else None
            cursor.fetchall.return_value = rows or []
            db.cursor.return_value = cursor
            return db

        rows = [{"role": "user", "content": "old context " * 100} for _ in range(9)]
        rows.append({"role": "user", "content": "current question"})
        completion = {"choices": [{"message": {"content": "READY"}}]}
        with patch.object(live, "get_db", side_effect=[fake_db(new_id=7), fake_db(rows=rows), fake_db()]), patch.object(
            live, "request_ai", return_value=completion
        ) as request_ai:
            response = self.client.post(
                "/api/chat",
                auth=self.auth,
                headers={"Origin": "http://localhost"},
                data={"message": "A" * live.MAX_CHAT_MESSAGE_CHARS, "mode": "study"},
            )

        self.assertEqual(response.status_code, 200)
        payload = request_ai.call_args.args[0]
        sent_history = payload["messages"][1:]
        self.assertEqual(sent_history[-1]["content"], "A" * live.MAX_CHAT_MESSAGE_CHARS)
        self.assertLessEqual(sum(len(item["content"]) for item in sent_history[:-1]), live.MAX_CHAT_HISTORY_CHARS)


if __name__ == "__main__":
    unittest.main()
