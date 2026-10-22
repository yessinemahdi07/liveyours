"""Bounded document extraction and study-context selection.

These functions keep file parsing limits explicit and independent from Flask
request handling so uploads can be tested at the service boundary.
"""
from dataclasses import dataclass
import csv
import io
import re
import os
import zipfile
from datetime import datetime

from PIL import Image
from openpyxl import load_workbook
from pypdf import PdfReader
import pytesseract

if os.name == "nt":
    pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
import xlrd


@dataclass(frozen=True)
class DocumentLimits:
    upload_bytes: int = 20 * 1024 * 1024
    document_pages: int = 50
    stored_document_chars: int = 30_000
    spreadsheet_sheets: int = 12
    spreadsheet_rows: int = 5_000
    spreadsheet_columns: int = 64
    spreadsheet_uncompressed_bytes: int = 50 * 1024 * 1024
    image_pixels: int = 13_000_000


DEFAULT_LIMITS = DocumentLimits()

def _spreadsheet_cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def _append_bounded_row(lines, current_length, row, limit):
    rendered = "\t".join(_spreadsheet_cell_text(cell)[:2_000] for cell in row).rstrip()
    if not rendered:
        return current_length, False
    prefix = "\n" if current_length else ""
    remaining = limit - current_length - len(prefix)
    if remaining <= 0:
        return current_length, True
    lines.append(prefix + rendered[:remaining])
    length = current_length + len(prefix) + min(len(rendered), remaining)
    return length, len(rendered) > remaining


def _extract_csv(raw, limits):
    try:
        source = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        source = raw.decode("cp1252")
    try:
        csv.field_size_limit(limits.upload_bytes)
        reader = csv.reader(io.StringIO(source, newline=""), strict=True)
        lines = []
        length = 0
        truncated = False
        for row_number, row in enumerate(reader, start=1):
            if row_number > limits.spreadsheet_rows:
                truncated = True
                break
            if len(row) > limits.spreadsheet_columns:
                row = row[:limits.spreadsheet_columns]
                truncated = True
            length, row_truncated = _append_bounded_row(lines, length, row, limits.stored_document_chars + 1)
            truncated = truncated or row_truncated
            if length > limits.stored_document_chars:
                truncated = True
                break
    except csv.Error as exc:
        raise ValueError("Could not read this CSV. Check that it is a valid, unencrypted text file.") from exc
    return "".join(lines)[:limits.stored_document_chars], truncated or length > limits.stored_document_chars


def _extract_xlsx(raw, limits):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 2_000 or sum(item.file_size for item in members) > limits.spreadsheet_uncompressed_bytes:
                raise ValueError("This workbook expands beyond the safe processing limit.")
            if any(item.file_size > limits.spreadsheet_uncompressed_bytes for item in members):
                raise ValueError("This workbook contains an oversized internal file.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Could not read this Excel workbook. Check that it is a valid .xlsx file.") from exc

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError("Could not read this Excel workbook. Check that it is a valid, unencrypted .xlsx file.") from exc

    lines = []
    length = 0
    truncated = False
    try:
        if len(workbook.worksheets) > limits.spreadsheet_sheets:
            truncated = True
        for sheet in workbook.worksheets[:limits.spreadsheet_sheets]:
            if length >= limits.stored_document_chars:
                truncated = True
                break
            header = f"Sheet: {sheet.title[:120]}"
            prefix = "\n" if length else ""
            lines.append(prefix + header)
            length += len(prefix) + len(header)
            max_row = min(sheet.max_row or limits.spreadsheet_rows, limits.spreadsheet_rows)
            max_column = min(sheet.max_column or limits.spreadsheet_columns, limits.spreadsheet_columns)
            if (sheet.max_row or 0) > limits.spreadsheet_rows or (sheet.max_column or 0) > limits.spreadsheet_columns:
                truncated = True
            for row_number, row in enumerate(
                sheet.iter_rows(max_row=max_row, max_col=max_column, values_only=True), start=1
            ):
                length, row_truncated = _append_bounded_row(lines, length, row, limits.stored_document_chars + 1)
                if row_truncated or length > limits.stored_document_chars:
                    truncated = True
                    break
            if length > limits.stored_document_chars:
                break
    finally:
        workbook.close()
    return "".join(lines)[:limits.stored_document_chars], truncated or length > limits.stored_document_chars


def _extract_xls(raw, limits):
    try:
        workbook = xlrd.open_workbook(file_contents=raw, on_demand=True)
    except Exception as exc:
        raise ValueError("Could not read this Excel workbook. Check that it is a valid, unencrypted .xls file.") from exc
    lines = []
    length = 0
    truncated = workbook.nsheets > limits.spreadsheet_sheets
    try:
        for sheet_index in range(min(workbook.nsheets, limits.spreadsheet_sheets)):
            sheet = workbook.sheet_by_index(sheet_index)
            if length >= limits.stored_document_chars:
                truncated = True
                break
            header = f"Sheet: {sheet.name[:120]}"
            prefix = "\n" if length else ""
            lines.append(prefix + header)
            length += len(prefix) + len(header)
            if sheet.nrows > limits.spreadsheet_rows or sheet.ncols > limits.spreadsheet_columns:
                truncated = True
            for row_index in range(min(sheet.nrows, limits.spreadsheet_rows)):
                row = [sheet.cell_value(row_index, col_index) for col_index in range(min(sheet.ncols, limits.spreadsheet_columns))]
                length, row_truncated = _append_bounded_row(lines, length, row, limits.stored_document_chars + 1)
                if row_truncated or length > limits.stored_document_chars:
                    truncated = True
                    break
            if length > limits.stored_document_chars:
                break
    finally:
        workbook.release_resources()
    return "".join(lines)[:limits.stored_document_chars], truncated or length > limits.stored_document_chars


def extract_document_text(upload, filename, limits=DEFAULT_LIMITS):
    """Extract bounded text from supported uploads without loading huge images into OCR."""
    extension = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if extension not in {'pdf', 'png', 'jpg', 'jpeg', 'csv', 'xlsx', 'xls'}:
        raise ValueError("Use a PDF, PNG, JPG, CSV, or Excel workbook (.xlsx or .xls).")

    raw = upload.read()
    if not raw:
        raise ValueError("The file is empty.")
    if len(raw) > limits.upload_bytes:
        raise ValueError("The file exceeds the 20 MB upload limit.")

    if extension == 'csv':
        return _extract_csv(raw, limits)
    if extension == 'xlsx':
        return _extract_xlsx(raw, limits)
    if extension == 'xls':
        return _extract_xls(raw, limits)

    if extension == 'pdf':
        reader = PdfReader(io.BytesIO(raw), strict=False)
        if len(reader.pages) > limits.document_pages:
            raise ValueError("Each PDF can contain at most 50 pages.")
        pieces = []
        length = 0
        for page in reader.pages:
            remaining = limits.stored_document_chars + 1 - length
            if remaining <= 0:
                break
            page_text = (page.extract_text() or "").strip()
            if page_text:
                piece = page_text[:remaining]
                pieces.append(piece)
                length += len(piece)
        text = "\n\n".join(pieces)
    else:
        image = Image.open(io.BytesIO(raw))
        if image.format not in {'PNG', 'JPEG'}:
            raise ValueError("Image content must match a PNG or JPEG file.")
        if image.width * image.height > limits.image_pixels:
            raise ValueError("Images must be 13 megapixels or smaller for safe OCR.")
        image.load()
        text = pytesseract.image_to_string(image)

    text = text.strip()
    truncated = len(text) > limits.stored_document_chars
    return text[:limits.stored_document_chars], truncated


def select_document_excerpt(text, query, max_chars):
    """Pick query-relevant text windows under a strict prompt budget."""
    if len(text) <= max_chars:
        return text
    chunk_size = 1_800
    chunks = [text[start:start + chunk_size] for start in range(0, len(text), chunk_size)]
    terms = {term.lower() for term in re.findall(r"[\w'-]{4,}", query)}
    scored = []
    for index, chunk in enumerate(chunks):
        words = set(term.lower() for term in re.findall(r"[\w'-]{4,}", chunk))
        scored.append((len(terms & words), index, chunk))
    separator = "\n\n[…sections omitted to keep the AI request small…]\n\n"
    opening = chunks[0][:min(500, max_chars // 4)]
    selected = [(0, opening)]
    remaining = max_chars - len(opening)
    ranked = sorted(scored[1:], key=lambda item: (-item[0], item[1]))
    for score, index, chunk in ranked:
        if remaining <= len(separator):
            break
        if score <= 0 and len(selected) > 1:
            continue
        excerpt = chunk[:remaining - len(separator)]
        if excerpt:
            selected.append((index, excerpt))
            remaining -= len(separator) + len(excerpt)
    if len(selected) == 1 and remaining > len(separator):
        index = min(1, len(chunks) - 1)
        selected.append((index, chunks[index][:remaining - len(separator)]))
    return separator.join(chunk for _, chunk in sorted(selected))[:max_chars]


def sample_document_text(text, max_chars):
    """Sample evenly across a document for bounded study generation."""
    if len(text) <= max_chars:
        return text, False
    window_count = max(1, max_chars // 2_000)
    separator = "\n\n[…sampled sections across the document…]\n\n"
    window_size = max(1, (max_chars - (window_count - 1) * len(separator)) // window_count)
    max_start = max(0, len(text) - window_size)
    starts = [round(i * max_start / max(1, window_count - 1)) for i in range(window_count)]
    excerpts = [text[start:start + window_size] for start in starts]
    return separator.join(excerpts)[:max_chars], True
